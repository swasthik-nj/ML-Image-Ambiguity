"""Generate the final end-to-end workflow documentation.

Run:

    python scripts/generate_final_workflow_doc.py

Produces: myDocs/Final_Workflow_and_Tools.docx
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = PROJECT_ROOT / "myDocs" / "Final_Workflow_and_Tools.docx"

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
DARK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x55, 0x55, 0x55)


def set_base_font(document: Document, name: str = "Calibri") -> None:
    style = document.styles["Normal"]
    style.font.name = name
    style.font.size = Pt(11)
    rpr = style.element.get_or_add_rPr()
    rFonts = rpr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)


def add_title_page(document: Document) -> None:
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Final System Workflow")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = ACCENT

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = subtitle.add_run(
        "Complete Tools, Implementation, and End-to-End Working Guide"
    )
    run2.italic = True
    run2.font.size = Pt(14)
    run2.font.color.rgb = DARK

    document.add_paragraph()
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(
        "Project: Explainable Image Ambiguity Prediction Using Human and "
        "AI-Generated Caption Diversity with Computer Vision Features\n"
        "Stack: COCO · Sentence-BERT · BLIP · OpenCV · Random Forest / XGBoost · "
        "SHAP · FastAPI · React (Ambiguity Lens)\n"
        "Document type: Final workflow & tools reference"
    )
    meta_run.font.size = Pt(11)
    meta_run.font.color.rgb = MUTED
    document.add_page_break()


def add_heading(document: Document, text: str, level: int = 1) -> None:
    heading = document.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = ACCENT


def add_body(document: Document, text: str) -> None:
    p = document.add_paragraph(text)
    p.paragraph_format.space_after = Pt(8)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        p = document.add_paragraph(item, style="List Bullet")
        p.paragraph_format.space_after = Pt(3)


def add_numbered(document: Document, items: list[str]) -> None:
    for item in items:
        p = document.add_paragraph(item, style="List Number")
        p.paragraph_format.space_after = Pt(3)


def add_code(document: Document, code: str) -> None:
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    for i, line in enumerate(code.strip("\n").split("\n")):
        run = p.add_run(("\n" if i > 0 else "") + line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x0A, 0x0A, 0x0A)
    p.paragraph_format.space_after = Pt(10)


def build() -> Path:
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    set_base_font(document)
    add_title_page(document)

    # ------------------------------------------------------------------
    add_heading(document, "1. Project Goal")
    add_body(
        document,
        "This system predicts whether an image is Low, Medium, or High "
        "ambiguity based on how much captions about that image disagree. "
        "Caption disagreement is measured with Sentence-BERT embeddings. "
        "Classical OpenCV features describe the visual appearance. A "
        "supervised classifier maps both feature groups to an ambiguity "
        "label. SHAP explains which features drove the decision. A FastAPI "
        "backend and React UI (Ambiguity Lens) make the pipeline interactive.",
    )
    add_body(
        document,
        "Important definition: Ambiguity here means caption disagreement, "
        "not visual clutter. A busy collage can still be Medium if captions "
        "agree on the content.",
    )

    # ------------------------------------------------------------------
    add_heading(document, "2. End-to-End Workflow (Big Picture)")
    add_body(document, "Offline research / training path:")
    add_numbered(
        document,
        [
            "Load COCO val2017 images + human captions.",
            "Embed captions with Sentence-BERT; compute caption diversity.",
            "Extract OpenCV features from each image.",
            "Build human_dataset.csv (and optionally ai_dataset.csv via BLIP).",
            "Assign Low / Medium / High labels from diversity thresholds.",
            "Enrich High-diversity samples if the High class is rare.",
            "Train Random Forest and XGBoost with class balancing; save models.",
            "Evaluate metrics and inspect SHAP explanations offline if needed.",
        ],
    )
    add_body(document, "Online prediction path (UI / API):")
    add_numbered(
        document,
        [
            "User uploads an image in Ambiguity Lens.",
            "If the filename is a COCO id (e.g. 000000538236.jpg), load human "
            "COCO captions automatically (unless BLIP is forced).",
            "Otherwise use user captions or generate diverse BLIP captions.",
            "Compute Sentence-BERT diversity + OpenCV features.",
            "Run the trained classifier → Low / Medium / High + probabilities.",
            "Compute SHAP feature contributions and return charts to the UI.",
            "Comparison page shows mean human vs AI caption diversity from CSVs.",
        ],
    )

    add_heading(document, "2.1 Workflow diagram (text)", level=2)
    add_code(
        document,
        """
COCO images + captions
        |
        v
[CocoDatasetLoader] --> captions per image
        |
        +--> [SentenceEmbeddingGenerator] --> embeddings
        |              |
        |              v
        |     [CaptionDiversityAnalyzer] --> diversity metrics
        |
        +--> [OpenCVFeatureExtractor] --> edge/entropy/brightness/...
        |
        v
[MLDatasetBuilder] --> human_dataset.csv / ai_dataset.csv
        |
        v
[AmbiguityLabelGenerator] --> Low / Medium / High
        |
        v
[enrich_high_diversity] (optional) --> more High rows
        |
        v
[ModelTrainer] --> best_model.joblib (+ RF / XGBoost)
        |
        v
[FastAPI Inference] <--> [React Ambiguity Lens]
        |
        +--> BLIP (optional AI captions)
        +--> SHAP explainer
        +--> /compare (human vs AI diversity)
""",
    )

    # ------------------------------------------------------------------
    add_heading(document, "3. Label Rules (Core Thesis Rule)")
    add_bullets(
        document,
        [
            "Low: caption_diversity < 0.35",
            "Medium: 0.35 ≤ caption_diversity < 0.65",
            "High: caption_diversity ≥ 0.65",
        ],
    )
    add_body(
        document,
        "These thresholds are kept intentionally. High still means captions "
        "strongly disagree. Training imbalance was fixed by adding more High "
        "images and oversampling during training — not by lowering the High "
        "threshold.",
    )

    # ------------------------------------------------------------------
    add_heading(document, "4. Architecture Layers")
    add_bullets(
        document,
        [
            "Domain / ML core — src/image_ambiguity/ "
            "(data, features, models, explainability, pipeline)",
            "Runnable scripts — src/*.py "
            "(create_dataset, generate_labels, train, enrich, BLIP AI set)",
            "Config — configs/default.yaml, .env, image_ambiguity.config",
            "API — backend/ + app.py (FastAPI)",
            "UI — frontend/ (Vite + React + TypeScript + Tailwind + Recharts)",
            "Artifacts — models/, results/, dataset/, myDocs/",
        ],
    )

    # ------------------------------------------------------------------
    add_heading(document, "5. Tools & Modules — How Each Works")

    add_heading(document, "5.1 CocoDatasetLoader", level=2)
    add_body(
        document,
        "File: src/image_ambiguity/data/coco_loader.py. "
        "Uses pycocotools to index captions_val2017.json and load images from "
        "dataset/val2017/. Provides get_captions(image_id), get_image(image_id), "
        "and random sampling. This is the data entry point for human captions.",
    )
    add_code(document, "python -c \"from image_ambiguity.data.coco_loader import CocoDatasetLoader\"")

    add_heading(document, "5.2 SentenceEmbeddingGenerator (Sentence-BERT)", level=2)
    add_body(
        document,
        "File: src/image_ambiguity/features/sentence_embeddings.py. "
        "Loads a Sentence-BERT model (configured in settings) and converts each "
        "caption into a dense vector. Similar captions → similar vectors. "
        "These embeddings are the input to diversity scoring.",
    )

    add_heading(document, "5.3 CaptionDiversityAnalyzer", level=2)
    add_body(
        document,
        "File: src/image_ambiguity/features/caption_diversity.py. "
        "Computes pairwise cosine similarities between caption embeddings for "
        "one image, then derives:",
    )
    add_bullets(
        document,
        [
            "average_similarity, minimum_similarity, maximum_similarity, std_similarity",
            "caption_diversity ≈ 1 − average_similarity "
            "(higher = captions disagree more)",
            "n_captions, n_pairs",
        ],
    )

    add_heading(document, "5.4 OpenCVFeatureExtractor", level=2)
    add_body(
        document,
        "File: src/image_ambiguity/features/cv_features.py. "
        "Reads the image pixels and extracts classical visual descriptors used "
        "by the classifier:",
    )
    add_bullets(
        document,
        [
            "edge_density — amount of edge structure (Canny-style density)",
            "entropy — information / texture complexity",
            "brightness — mean intensity",
            "contrast — intensity spread",
            "color_variance — color channel variation",
            "texture — local texture strength",
        ],
    )

    add_heading(document, "5.5 BlipCaptionGenerator", level=2)
    add_body(
        document,
        "File: src/image_ambiguity/features/blip_captions.py. "
        "Uses Salesforce/blip-image-captioning-base to generate AI captions with "
        "multiple decoding strategies: beam search, top-k sampling, nucleus "
        "(top-p) sampling, plus prompt-conditioned generations for more variety. "
        "Used for ai_dataset.csv and for UI predictions when BLIP is forced.",
    )
    add_body(
        document,
        "Note: BLIP captions often agree more than human COCO captions. The same "
        "image can be High under humans and Medium under BLIP. That is a research "
        "finding (human vs AI caption agreement), not a UI bug.",
    )

    add_heading(document, "5.6 MLDatasetBuilder", level=2)
    add_body(
        document,
        "File: src/image_ambiguity/pipeline/dataset_builder.py. "
        "Joins caption-diversity metrics and OpenCV features into one row per "
        "image. Can inject alternate captions (e.g. BLIP) via captions_by_image. "
        "Output columns include image_id, diversity features, and CV features.",
    )

    add_heading(document, "5.7 AmbiguityLabelGenerator", level=2)
    add_body(
        document,
        "File: src/image_ambiguity/pipeline/label_generator.py. "
        "Applies the Low / Medium / High rules to caption_diversity and writes "
        "ambiguity_label. Also supports distribution summaries and plots.",
    )

    add_heading(document, "5.8 ModelTrainer", level=2)
    add_body(
        document,
        "File: src/image_ambiguity/models/trainer.py. "
        "Trains RandomForestClassifier and XGBoostClassifier on the 11 numeric "
        "features to predict ambiguity_label. Pipeline steps:",
    )
    add_numbered(
        document,
        [
            "Load labeled CSV; encode labels as 0=Low, 1=Medium, 2=High.",
            "Stratified 80/20 train/test split.",
            "Optional oversampling of minority classes on the train set only "
            "(balance_classes=True by default).",
            "Baseline k-fold CV + GridSearchCV hyperparameter tuning.",
            "Evaluate Accuracy / Precision / Recall / F1 / ROC-AUC on test set.",
            "Save best model with joblib.",
        ],
    )
    add_body(
        document,
        "Feature columns: average_similarity, minimum_similarity, "
        "maximum_similarity, std_similarity, caption_diversity, edge_density, "
        "entropy, brightness, contrast, color_variance, texture.",
    )

    add_heading(document, "5.9 SHAPExplainer", level=2)
    add_body(
        document,
        "File: src/image_ambiguity/explainability/shap_explainer.py. "
        "Uses SHAP to attribute the model prediction to individual features. "
        "The API /explain endpoint returns top contributions and a text summary "
        "shown in the UI.",
    )

    add_heading(document, "5.10 AmbiguityInferenceService (API brain)", level=2)
    add_body(
        document,
        "File: backend/app/services/inference.py. "
        "Orchestrates upload storage, caption resolution, feature extraction, "
        "prediction, and SHAP. Caption resolution order:",
    )
    add_numbered(
        document,
        [
            "If the user provides ≥2 captions → use them (source=user).",
            "Else if not force_blip and filename/upload maps to a COCO id → "
            "load human COCO captions (source=coco_human).",
            "Else generate BLIP captions (source=blip).",
        ],
    )

    add_heading(document, "5.11 Compare service", level=2)
    add_body(
        document,
        "File: backend/app/services/compare.py. "
        "Reads dataset/human_dataset.csv and dataset/ai_dataset.csv and returns "
        "mean caption diversity plus histograms for the Comparison page.",
    )

    # ------------------------------------------------------------------
    add_heading(document, "6. Runnable Scripts (CLI Tools)")

    rows = [
        (
            "src/create_dataset.py",
            "Build human_dataset.csv from sampled COCO images "
            "(diversity + OpenCV features).",
        ),
        (
            "src/generate_labels.py",
            "Apply Low/Medium/High thresholds to caption_diversity.",
        ),
        (
            "src/enrich_high_diversity.py",
            "Mine unused High-diversity COCO images and append them "
            "(keeps High ≥ 0.65).",
        ),
        (
            "src/train.py",
            "Train RF + XGBoost with optional class balancing; save models "
            "and metrics JSON.",
        ),
        (
            "src/create_ai_dataset.py",
            "Generate BLIP captions for the same image ids and build "
            "ai_dataset.csv.",
        ),
        (
            "src/blip_caption.py / features/blip_captions.py",
            "Standalone / library BLIP caption generation.",
        ),
        (
            "src/shap_analysis.py",
            "Offline SHAP analysis helper for trained models.",
        ),
    ]
    for path, desc in rows:
        add_body(document, f"{path} — {desc}")

    add_heading(document, "6.1 Typical offline command sequence", level=2)
    add_code(
        document,
        """
# from project root, with venv active
$env:PYTHONPATH="src;."

python src/create_dataset.py
python src/generate_labels.py
python src/enrich_high_diversity.py --max-high 40 --high-min 0.65
python src/train.py --balance-classes

# optional AI comparison set (slow on CPU)
python src/create_ai_dataset.py --device cpu
""",
    )

    # ------------------------------------------------------------------
    add_heading(document, "7. FastAPI Backend")
    add_body(
        document,
        "Entrypoint: app.py → uvicorn app:app --reload. "
        "Main app: backend/app/main.py. Routes: backend/app/api/routes.py.",
    )
    add_bullets(
        document,
        [
            "POST /upload — store image; return upload_id; attach COCO captions "
            "if filename matches a COCO id",
            "POST /features — OpenCV + diversity features",
            "POST /predict — ambiguity class + probabilities",
            "POST /explain — prediction + SHAP explanation",
            "GET /compare — human vs AI mean diversity + histograms",
            "GET /health — health check",
        ],
    )
    add_code(
        document,
        """
$env:PYTHONPATH="src;."
uvicorn app:app --reload
# Swagger: http://127.0.0.1:8000/docs
""",
    )

    # ------------------------------------------------------------------
    add_heading(document, "8. React Frontend — Ambiguity Lens")
    add_body(
        document,
        "Location: frontend/. Stack: Vite, React, TypeScript, Tailwind CSS v4, "
        "Recharts, React Router. Proxies /api to the FastAPI backend.",
    )
    add_bullets(
        document,
        [
            "Home — brand landing (Ambiguity Lens)",
            "Prediction — upload image, captions or BLIP, show label / "
            "confidence / diversity / SHAP / OpenCV charts",
            "Comparison — human vs AI caption diversity bars and histograms",
            "About — stack and run notes",
        ],
    )
    add_code(
        document,
        """
cd frontend
npm install
npm run dev
# UI: http://localhost:5173
""",
    )
    add_body(
        document,
        "For COCO val filenames (000000######.jpg), the UI turns BLIP off by "
        "default and loads human captions so High images stay High. Check BLIP "
        "only when you intentionally want AI captions.",
    )

    # ------------------------------------------------------------------
    add_heading(document, "9. Artifacts & Outputs")
    add_bullets(
        document,
        [
            "dataset/human_dataset.csv — labeled human-caption feature table",
            "dataset/ai_dataset.csv — BLIP-based feature table (optional)",
            "dataset/val2017/, dataset/annotations/ — COCO local data",
            "models/best_model.joblib — production classifier for the API",
            "models/random_forest.joblib, models/xgboost.joblib",
            "results/metrics/training_metrics.json",
            "results/uploads/ — uploaded demo images + metadata",
            "results/captions/ — cached BLIP outputs when used",
            "myDocs/*.docx — generated explanation documents",
        ],
    )

    # ------------------------------------------------------------------
    add_heading(document, "10. Feature Vector Used by the Model")
    add_body(
        document,
        "Eleven numeric inputs (same order in training and inference):",
    )
    add_numbered(
        document,
        [
            "average_similarity",
            "minimum_similarity",
            "maximum_similarity",
            "std_similarity",
            "caption_diversity",
            "edge_density",
            "entropy",
            "brightness",
            "contrast",
            "color_variance",
            "texture",
        ],
    )

    # ------------------------------------------------------------------
    add_heading(document, "11. Recent Improvements Included in This Build")
    add_bullets(
        document,
        [
            "High-class enrichment (High samples increased substantially).",
            "Train-split oversampling + Random Forest class_weight=balanced.",
            "Auto COCO human caption lookup for dataset images in the API/UI.",
            "force_blip flag when pure AI captions are required.",
            "Richer BLIP decoding (temperature, more sequences, prompt variants).",
            "caption_source returned in API responses (user / coco_human / blip).",
        ],
    )

    # ------------------------------------------------------------------
    add_heading(document, "12. How to Demo Correctly")
    add_numbered(
        document,
        [
            "Start API with PYTHONPATH including src.",
            "Start frontend with npm run dev.",
            "For High demo: upload dataset/val2017/000000538236.jpg "
            "(leave BLIP unchecked).",
            "Expect High with diversity around 0.65–0.80 and source "
            "COCO human captions.",
            "Optionally enable BLIP on the same image to show AI captions "
            "may score Medium — useful for the Comparison thesis.",
            "Open Comparison page after ai_dataset.csv exists to contrast "
            "mean human vs AI diversity.",
        ],
    )

    # ------------------------------------------------------------------
    add_heading(document, "13. Suggested High-Ambiguity Test Images")
    add_bullets(
        document,
        [
            "dataset/val2017/000000538236.jpg — diversity ≈ 0.80",
            "dataset/val2017/000000546823.jpg — diversity ≈ 0.79",
            "dataset/val2017/000000323263.jpg — diversity ≈ 0.79",
            "dataset/val2017/000000263425.jpg — diversity ≈ 0.79",
            "dataset/val2017/000000060932.jpg — diversity ≈ 0.76",
        ],
    )

    # ------------------------------------------------------------------
    add_heading(document, "14. Quick Troubleshooting")
    add_bullets(
        document,
        [
            "All predictions Medium with identical captions → you reused the "
            "same manual captions for every image; use COCO captions or BLIP "
            "per image.",
            "COCO High image + BLIP → may be Medium because AI captions agree; "
            "uncheck BLIP for human-High demos.",
            "npm ENOENT package.json → frontend source missing; restore "
            "frontend files then npm run dev.",
            "API model load errors → run python src/train.py; ensure "
            "models/best_model.joblib exists; install xgboost if needed.",
            "Slow BLIP on CPU → expected; use human captions for faster demos.",
        ],
    )

    # ------------------------------------------------------------------
    add_heading(document, "15. Research Takeaway")
    add_body(
        document,
        "The system operationalizes image ambiguity as measurable caption "
        "disagreement, enriched with visual OpenCV cues, predicted by classical "
        "ML, and explained with SHAP. Human and AI caption pipelines share the "
        "same diversity math, enabling a direct comparison. Keeping High ≥ 0.65 "
        "preserves a strong thesis definition while enrichment and balancing "
        "make training fair. The interactive Ambiguity Lens app demonstrates "
        "the full research pipeline from image to explained prediction.",
    )

    document.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
