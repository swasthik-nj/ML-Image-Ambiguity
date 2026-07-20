"""Generate the FastAPI backend explanation document.

Run:

    python scripts/generate_fastapi_doc.py

Produces: myDocs/FastAPI_Backend_Explanation.docx
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = PROJECT_ROOT / "myDocs" / "FastAPI_Backend_Explanation.docx"

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
DARK = RGBColor(0x22, 0x22, 0x22)


def set_base_font(document: Document, name: str = "Calibri") -> None:
    style = document.styles["Normal"]
    style.font.name = name
    style.font.size = Pt(11)
    rpr = style.element.get_or_add_rPr()
    rFonts = rpr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)


def add_title_page(document: Document) -> None:
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("FastAPI Backend")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = ACCENT

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = subtitle.add_run(
        "REST API for Image Upload, Feature Extraction, "
        "Ambiguity Prediction, and SHAP Explanation"
    )
    run2.italic = True
    run2.font.size = Pt(14)
    run2.font.color.rgb = DARK

    document.add_paragraph()
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(
        "Project: Explainable Image Ambiguity Prediction Using Human and "
        "AI-Generated Caption Diversity with Computer Vision Features\n"
        "Entrypoint: app.py → backend.app.main:app\n"
        "Run: uvicorn app:app --reload"
    )
    meta_run.font.size = Pt(11)
    meta_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    document.add_page_break()


def add_heading(document: Document, text: str, level: int = 1) -> None:
    heading = document.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = ACCENT


def add_body(document: Document, text: str) -> None:
    p = document.add_paragraph(text)
    p.paragraph_format.space_after = Pt(8)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_numbered(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Number")


def add_code(document: Document, code: str) -> None:
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    for i, line in enumerate(code.strip("\n").split("\n")):
        run = p.add_run(("\n" if i > 0 else "") + line)
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x0A, 0x0A, 0x0A)
    p.paragraph_format.space_after = Pt(10)


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    document.add_paragraph()


def build() -> None:
    document = Document()
    set_base_font(document)
    add_title_page(document)

    add_heading(document, "1. Purpose")
    add_body(
        document,
        "The FastAPI backend exposes the research pipeline as a REST API. "
        "A client can upload an image, extract caption-diversity and OpenCV "
        "features, predict Low/Medium/High ambiguity, and inspect a SHAP "
        "explanation of that prediction.",
    )

    add_heading(document, "2. Tools Used")
    add_table(
        document,
        ["Tool", "Role"],
        [
            ["FastAPI", "REST API framework and OpenAPI docs"],
            ["Uvicorn", "ASGI server (`uvicorn app:app --reload`)"],
            ["Pydantic", "Request/response schemas"],
            ["Pillow", "Validate and inspect uploaded images"],
            ["OpenCV", "Classical image features"],
            ["Sentence-BERT", "Caption embeddings for diversity"],
            ["BLIP + Transformers + PyTorch", "Optional AI caption generation"],
            ["scikit-learn / XGBoost + joblib", "Load trained ambiguity classifier"],
            ["SHAP", "Feature attribution explanations"],
        ],
    )

    add_heading(document, "3. Project Layout")
    add_table(
        document,
        ["Path", "Responsibility"],
        [
            ["app.py", "Root entrypoint for `uvicorn app:app`"],
            ["backend/app/main.py", "FastAPI app, CORS, /health, /"],
            ["backend/app/api/routes.py", "REST endpoints"],
            ["backend/app/schemas.py", "Pydantic response models"],
            [
                "backend/app/services/inference.py",
                "Upload handling + feature/predict/explain logic",
            ],
            ["models/best_model.joblib", "Trained classifier artifact"],
            ["results/uploads/", "Saved uploaded images"],
        ],
    )

    add_heading(document, "4. REST Endpoints")
    add_table(
        document,
        ["Method", "Path", "Purpose"],
        [
            ["GET", "/health", "Liveness check"],
            ["GET", "/", "API metadata / endpoint index"],
            ["POST", "/upload", "Store image; return upload_id"],
            ["POST", "/features", "OpenCV + caption-diversity features"],
            ["POST", "/predict", "Predicted ambiguity + features"],
            ["POST", "/explain", "Prediction + SHAP explanation"],
        ],
    )
    add_body(
        document,
        "Interactive docs are available at /docs (Swagger UI) and /redoc.",
    )

    add_heading(document, "5. How Inference Works")
    add_numbered(
        document,
        [
            "Accept an uploaded image (or a previous upload_id).",
            "Extract OpenCV features from the image pixels.",
            "Obtain captions: either a JSON captions form field, or "
            "generate them with BLIP.",
            "Embed captions with Sentence-BERT and compute diversity "
            "(Diversity = 1 - Average Similarity).",
            "Build the 11-feature vector expected by the trained model.",
            "Predict Low / Medium / High ambiguity with probabilities.",
            "For /explain, compute SHAP contributions for the predicted class.",
        ],
    )

    add_heading(document, "6. Request Fields (Swagger / multipart form)")
    add_table(
        document,
        ["Field", "Required?", "Example / notes"],
        [
            ["file", "One of file or upload_id", "Choose an image in Swagger"],
            ["upload_id", "One of file or upload_id", "From POST /upload response"],
            [
                "captions",
                "Optional",
                '["a kitchen scene", "a person cooking"] '
                "(at least 2). Leave empty to use BLIP.",
            ],
            ["top_n", "/explain only", "Default 5"],
        ],
    )
    add_body(
        document,
        "Important: delete Swagger’s default placeholder text `string` "
        "from optional fields. For captions, provide valid JSON or clear "
        "the field completely.",
    )

    add_heading(document, "7. Response Contents")
    add_bullets(
        document,
        [
            "predicted_ambiguity: Low / Medium / High",
            "confidence and per-class probabilities",
            "caption_diversity metrics (average/min/max/std + score)",
            "opencv_features (edge_density, entropy, brightness, contrast, "
            "color_variance, texture)",
            "captions used for the request",
            "shap_explanation + summary text on /explain",
        ],
    )

    add_heading(document, "8. Usage")
    add_heading(document, "8.1 Start the server", level=2)
    add_code(
        document,
        """
# from project root
$env:PYTHONPATH="src;."
uvicorn app:app --reload
""",
    )
    add_body(document, "Then open: http://127.0.0.1:8000/docs")

    add_heading(document, "8.2 Recommended Swagger flow", level=2)
    add_numbered(
        document,
        [
            "POST /upload → select an image → copy upload_id from the response.",
            "POST /predict → paste upload_id; set captions to a JSON array "
            "of at least two strings (faster than BLIP).",
            "POST /explain → same upload_id + captions to inspect SHAP.",
        ],
    )

    add_heading(document, "8.3 Example captions value", level=2)
    add_code(
        document,
        """
["a person in a kitchen", "someone cooking near a stove"]
""",
    )

    add_heading(document, "8.4 Example curl", level=2)
    add_code(
        document,
        """
curl -X POST "http://127.0.0.1:8000/predict" ^
  -H "accept: application/json" ^
  -F "upload_id=YOUR_UPLOAD_ID" ^
  -F "captions=[\"a person in a kitchen\", \"someone cooking near a stove\"]"
""",
    )

    add_heading(document, "9. Dependencies / Runtime Notes")
    add_bullets(
        document,
        [
            "Install API stack: pip install fastapi \"uvicorn[standard]\"",
            "Install ML stack from requirements.txt "
            "(torch, transformers, sentence-transformers, opencv-python, "
            "scikit-learn, xgboost, shap, joblib).",
            "best_model.joblib may require xgboost; the service can fall "
            "back to random_forest.joblib if needed.",
            "First BLIP/Sentence-BERT use downloads model weights "
            "(slower on CPU).",
        ],
    )

    add_heading(document, "10. REST Practices Followed")
    add_bullets(
        document,
        [
            "Resource-oriented endpoints with clear verbs via HTTP methods",
            "Consistent JSON error bodies with appropriate status codes "
            "(400 / 404 / 422 / 500)",
            "OpenAPI documentation via /docs",
            "Separation of routes (api/) and business logic (services/)",
            "Typed response models with Pydantic",
            "CORS configured from project settings",
        ],
    )

    add_heading(document, "11. Testing")
    add_code(
        document,
        """
python -m pytest tests/integration/test_api_health.py tests/integration/test_api_endpoints.py -q
""",
    )

    add_heading(document, "12. References")
    add_bullets(
        document,
        [
            "FastAPI documentation: https://fastapi.tiangolo.com/",
            "Uvicorn documentation: https://www.uvicorn.org/",
            "Lundberg, S. M., & Lee, S.-I. (2017). SHAP. NeurIPS 2017.",
            "Li, J. et al. (2022). BLIP. ICML 2022.",
        ],
    )

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")


if __name__ == "__main__":
    build()
