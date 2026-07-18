"""One-off script to generate the Training Pipeline explanation document.

Run once with:

    python scripts/generate_training_pipeline_doc.py

Produces: myDocs/Training_Pipeline_Explanation.docx
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = PROJECT_ROOT / "myDocs" / "Training_Pipeline_Explanation.docx"

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
DARK = RGBColor(0x22, 0x22, 0x22)


def set_base_font(document: Document, name: str = "Calibri") -> None:
    style = document.styles["Normal"]
    style.font.name = name
    style.font.size = Pt(11)
    rpr = style.element.get_or_add_rPr()
    rFonts = rpr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)


def add_title_page(document: Document) -> None:
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Training Classifiers to Predict Ambiguity")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = ACCENT

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = subtitle.add_run(
        "Random Forest and XGBoost with Cross-Validation and Hyperparameter Tuning"
    )
    run2.italic = True
    run2.font.size = Pt(15)
    run2.font.color.rgb = DARK

    document.add_paragraph()
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(
        "Project: Explainable Image Ambiguity Prediction Using Human and "
        "AI-Generated Caption Diversity with Computer Vision Features\n"
        "Module: image_ambiguity.models.trainer.ModelTrainer\n"
        "Output: models/best_model.joblib, results/metrics/training_metrics.json"
    )
    meta_run.font.size = Pt(11)
    meta_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    document.add_page_break()


def add_heading(document: Document, text: str, level: int = 1) -> None:
    heading = document.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = ACCENT


def add_body(document: Document, text: str) -> None:
    p = document.add_paragraph(text)
    p.paragraph_format.space_after = Pt(8)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_numbered(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Number")


def add_code(document: Document, code: str) -> None:
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    for i, line in enumerate(code.strip("\n").split("\n")):
        run = p.add_run(("\n" if i > 0 else "") + line)
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x0A, 0x0A, 0x0A)
    p.paragraph_format.space_after = Pt(10)


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    document.add_paragraph()


def build() -> None:
    document = Document()
    set_base_font(document)
    add_title_page(document)

    # 1. Motivation
    add_heading(document, "1. Motivation")
    add_body(
        document,
        "Every prior module in this project -- Sentence-BERT embeddings, "
        "caption diversity, OpenCV features, and rule-based ambiguity "
        "labels -- produces descriptive statistics or fixed thresholds, "
        "not a predictive model. The Training Pipeline is the first "
        "module that actually learns a mapping from image features to "
        "an ambiguity category, turning human_dataset.csv from a static "
        "feature table into a reusable, deployable classifier."
    )
    add_body(
        document,
        "Two complementary tree-ensemble algorithms are trained side by "
        "side -- Random Forest and XGBoost -- so their behavior can be "
        "directly compared on the same 80/20 split, the same "
        "cross-validation folds, and the same evaluation metrics before "
        "the stronger of the two is persisted for reuse."
    )

    # 2. What Random Forest is and why it is used
    add_heading(document, "2. Random Forest: What It Is and Why It Is Used")
    add_body(
        document,
        "A Random Forest is an ensemble of many decision trees, each "
        "trained on a bootstrap-resampled subset of the training rows "
        "and a random subset of the feature columns at every split. "
        "Its final prediction is a majority vote (classification) "
        "across all trees. This 'bagging' of many weak, decorrelated "
        "trees reduces the variance and overfitting that a single deep "
        "decision tree would suffer on a dataset this small."
    )
    add_bullets(
        document,
        [
            "Handles the 11 numeric features (caption-diversity + OpenCV "
            "statistics) directly, with no scaling or normalization "
            "required -- tree splits are invariant to monotonic feature "
            "transforms.",
            "Naturally exposes per-feature importances, which double as "
            "an explainability artifact: which of the 11 inputs the "
            "model actually relies on to separate Low/Medium/High.",
            "Robust to the mixed, differently-scaled feature ranges in "
            "this dataset (cosine similarities in [0, 1] alongside pixel "
            "statistics in the hundreds or thousands, e.g. color_variance).",
            "Serves as a well-understood, easy-to-explain baseline "
            "against which the more complex XGBoost model is judged.",
        ],
    )

    # 3. What XGBoost is and why it is used
    add_heading(document, "3. XGBoost: What It Is and Why It Is Used")
    add_body(
        document,
        "XGBoost (Extreme Gradient Boosting) is also an ensemble of "
        "decision trees, but built sequentially rather than in parallel: "
        "each new tree is trained to correct the residual errors of the "
        "ensemble built so far, using gradient descent on a "
        "differentiable loss (multi-class log-loss here). This typically "
        "yields higher accuracy than bagging on structured/tabular data, "
        "at the cost of being more sensitive to hyperparameters."
    )
    add_bullets(
        document,
        [
            "Boosting directly targets the hardest-to-classify rows "
            "(e.g. borderline Medium/High diversity scores), often "
            "outperforming Random Forest on tabular classification "
            "benchmarks -- confirmed in this project's own run, where "
            "XGBoost reached 100% test accuracy versus Random Forest's "
            "98.3%.",
            "Regularization terms (max_depth, learning_rate, "
            "n_estimators) give fine-grained control over overfitting, "
            "which matters on a 300-row dataset with a severely "
            "imbalanced High-ambiguity class (only 4 images).",
            "Same feature-importance interface as Random Forest, so the "
            "two models can be compared not just on accuracy but on "
            "which features they consider predictive.",
            "Widely used in production ML systems and Kaggle-style "
            "competitions for tabular data, making it a credible "
            "second model to report alongside a simpler baseline.",
        ],
    )

    # 4. Role of each in this project
    add_heading(document, "4. Role of Random Forest and XGBoost in This Project")
    add_body(
        document,
        "Both models are trained on the exact same input: the 11 "
        "feature columns already computed by earlier modules "
        "(average/min/max/std similarity, caption_diversity, "
        "edge_density, entropy, brightness, contrast, color_variance, "
        "texture), predicting the rule-based ambiguity_label "
        "(Low/Medium/High) produced by the labeling module. Neither "
        "model replaces the earlier feature-engineering steps -- they "
        "sit on top of them, learning which combinations of those "
        "features are jointly predictive of ambiguity, rather than "
        "relying on caption_diversity alone via a single fixed "
        "threshold."
    )
    add_body(
        document,
        "Concretely: Random Forest acts as the interpretable, "
        "low-variance baseline; XGBoost acts as the higher-capacity "
        "challenger. ModelTrainer.select_best() then automatically "
        "promotes whichever model scores higher on the chosen test-set "
        "metric (accuracy by default) to models/best_model.joblib, so "
        "the rest of the project (e.g. a future inference API) can load "
        "one artifact without caring which algorithm produced it."
    )

    # 5. Pipeline steps
    add_heading(document, "5. How the Pipeline Is Assembled")
    add_numbered(
        document,
        [
            "Prepare data: read human_dataset.csv, keep only rows with a "
            "valid Low/Medium/High label, encode labels as integers "
            "0/1/2, and median-impute any missing feature values.",
            "Split 80/20: a stratified train_test_split preserves the "
            "Low/Medium/High class proportions in both the training and "
            "test sets (falling back to a plain random split if a class "
            "is too small to stratify).",
            "Baseline cross-validation: each model, with default "
            "hyperparameters, is scored with stratified k-fold CV on the "
            "training set to establish a pre-tuning reference accuracy.",
            "Hyperparameter tuning: GridSearchCV re-runs cross-validation "
            "across a small parameter grid per model (tree count, depth, "
            "learning rate, etc.), selecting the combination with the "
            "best mean CV accuracy and refitting it on the full training "
            "set.",
            "Evaluate on the held-out test set: Accuracy, macro-averaged "
            "Precision/Recall/F1, and macro one-vs-rest ROC-AUC are "
            "computed once per model on data neither model has ever "
            "seen during tuning.",
            "Select and persist: the model with the higher chosen metric "
            "(accuracy by default) is saved as models/best_model.joblib; "
            "both individual models are also saved for comparison, "
            "alongside a JSON metrics report.",
        ],
    )
    add_body(
        document,
        "Crucially, the fold count for cross-validation and grid search "
        "is computed dynamically from the smallest class's sample count "
        "(_resolve_cv_folds), and hyperparameter tuning is skipped "
        "gracefully -- falling back to default parameters -- if a class "
        "has fewer than two samples. This prevents the pipeline from "
        "crashing on the severely imbalanced High-ambiguity class (only "
        "4 of 300 images)."
    )

    # 6. Hyperparameters tuned
    add_heading(document, "6. Hyperparameters Tuned")
    add_table(
        document,
        ["Model", "Grid Searched", "Best Found (this run)"],
        [
            [
                "Random Forest",
                "n_estimators in {100, 200}\nmax_depth in {None, 8, 16}\nmin_samples_leaf in {1, 2}",
                "n_estimators=100, max_depth=None, min_samples_leaf=1",
            ],
            [
                "XGBoost",
                "n_estimators in {100, 200}\nmax_depth in {3, 5}\nlearning_rate in {0.05, 0.1}",
                "n_estimators=100, max_depth=3, learning_rate=0.05",
            ],
        ],
    )

    # 7. Worked example
    add_heading(document, "7. Worked Example (from this codebase)")
    add_body(
        document,
        "Running python src/train.py against the 300-row "
        "human_dataset.csv (240 train / 60 test) produced:"
    )
    add_table(
        document,
        ["Metric", "Random Forest", "XGBoost"],
        [
            ["Baseline 3-fold CV Accuracy", "98.33%", "99.17%"],
            ["Tuned CV Score", "98.33%", "99.17%"],
            ["Test Accuracy", "98.33%", "100.00%"],
            ["Test Precision (macro)", "65.79%", "100.00%"],
            ["Test Recall (macro)", "66.67%", "100.00%"],
            ["Test F1 (macro)", "66.22%", "100.00%"],
            ["Test ROC-AUC (macro, OvR)", "100.00%", "100.00%"],
        ],
    )
    add_body(
        document,
        "XGBoost was automatically selected as the best model (highest "
        "test accuracy) and saved to models/best_model.joblib, alongside "
        "models/random_forest.joblib and models/xgboost.joblib for "
        "direct comparison, and results/metrics/training_metrics.json "
        "for the full numeric report."
    )

    # 8. Reading the precision/recall gap
    add_heading(document, "8. Why Random Forest's Precision/Recall Look Low Despite 98% Accuracy")
    add_body(
        document,
        "Accuracy and macro-averaged Precision/Recall/F1 answer "
        "different questions. Accuracy simply asks 'what fraction of "
        "the 60 test images were labeled correctly?' -- and with Low "
        "and Medium making up 296 of 300 images, a model can score "
        "very high on accuracy while still struggling on the rare High "
        "class."
    )
    add_body(
        document,
        "Macro-averaging computes Precision/Recall/F1 separately for "
        "each of the three classes and then averages them unweighted -- "
        "so the High class (only 4 images total, roughly 1 in the test "
        "split) counts exactly as much as Low or Medium. Random "
        "Forest's single miss on that lone High-ambiguity test image "
        "drags its macro Precision/Recall/F1 down to ~66%, even though "
        "its overall accuracy remained 98.3%. XGBoost happened to "
        "classify every test image correctly in this run, hence its "
        "perfect scores across the board."
    )
    add_body(
        document,
        "This is a direct, quantitative illustration of why relying on "
        "accuracy alone is misleading on imbalanced datasets -- and why "
        "this pipeline reports Precision, Recall, F1, and ROC-AUC "
        "alongside Accuracy rather than Accuracy in isolation."
    )

    # 9. A note on label provenance
    add_heading(document, "9. A Note on Label Provenance")
    add_body(
        document,
        "Because ambiguity_label is itself a deterministic threshold "
        "function of one of the eleven input features "
        "(caption_diversity), both models can, in principle, recover "
        "that exact rule from the data -- which is a major contributor "
        "to the very high accuracy seen here. For a benchmark that "
        "measures genuine generalization beyond the labeling rule "
        "itself, a natural next step is to retrain with caption_"
        "diversity excluded from the feature set, so the models must "
        "predict ambiguity purely from the remaining caption-agreement "
        "statistics and OpenCV image features."
    )

    # 10. Explainability
    add_heading(document, "10. Why This Design Supports Explainability")
    add_body(
        document,
        "Both Random Forest and XGBoost expose feature_importances_ "
        "out of the box, and both are natively supported by SHAP "
        "(already a project dependency), so any prediction from "
        "models/best_model.joblib can be decomposed into a per-feature "
        "contribution -- e.g. 'this image was labeled High-ambiguity "
        "primarily because of caption_diversity and edge_density.' "
        "Because every feature the models consume is itself already "
        "documented and traceable back to raw captions or pixels (via "
        "the earlier modules), the full chain from evidence to "
        "prediction remains interpretable end-to-end."
    )

    # 11. Talking points
    add_heading(document, "11. Conference Talking Points (Summary)")
    add_bullets(
        document,
        [
            "Two tree-ensemble classifiers -- Random Forest (bagging) "
            "and XGBoost (gradient boosting) -- are trained on the same "
            "11-feature table to predict Low/Medium/High ambiguity, with "
            "an 80/20 stratified split.",
            "Both baseline cross-validation and GridSearchCV "
            "hyperparameter tuning are performed per model, with fold "
            "counts adapted automatically to the smallest class size "
            "(only 4 High-ambiguity images).",
            "On this run, XGBoost reached 100% test accuracy versus "
            "Random Forest's 98.3%, consistent with boosting's known "
            "edge over bagging on tabular data.",
            "Reporting Accuracy alongside macro Precision/Recall/F1 and "
            "ROC-AUC exposed a real class-imbalance effect: Random "
            "Forest's macro scores (~66%) were far lower than its "
            "accuracy (98.3%) due to a single miss on the rare High "
            "class.",
            "The best model by accuracy is automatically selected and "
            "saved with joblib to models/best_model.joblib, decoupling "
            "model training from downstream deployment.",
            "Because ambiguity_label is rule-derived from caption_"
            "diversity, near-perfect scores partly reflect the labeling "
            "rule itself -- motivating a follow-up experiment without "
            "that feature to measure genuine predictive generalization.",
        ],
    )

    # 12. References
    add_heading(document, "12. References")
    add_bullets(
        document,
        [
            "Breiman, L. (2001). Random Forests. Machine Learning, "
            "45(1), 5-32.",
            "Chen, T., & Guestrin, C. (2016). XGBoost: A Scalable Tree "
            "Boosting System. KDD 2016.",
            "Pedregosa, F. et al. (2011). Scikit-learn: Machine Learning "
            "in Python. JMLR 12.",
            "Lundberg, S. M., & Lee, S.-I. (2017). A Unified Approach to "
            "Interpreting Model Predictions (SHAP). NeurIPS 2017.",
        ],
    )

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")


if __name__ == "__main__":
    build()
