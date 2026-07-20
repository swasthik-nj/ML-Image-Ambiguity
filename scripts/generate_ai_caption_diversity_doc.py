"""Generate the AI Caption Diversity explanation document.

Run:

    python scripts/generate_ai_caption_diversity_doc.py

Produces: myDocs/AI_Caption_Diversity_Explanation.docx
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = PROJECT_ROOT / "myDocs" / "AI_Caption_Diversity_Explanation.docx"

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
DARK = RGBColor(0x22, 0x22, 0x22)


def set_base_font(document: Document, name: str = "Calibri") -> None:
    style = document.styles["Normal"]
    style.font.name = name
    style.font.size = Pt(11)
    rpr = style.element.get_or_add_rPr()
    rFonts = rpr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)


def add_title_page(document: Document) -> None:
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("AI Caption Diversity Pipeline")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = ACCENT

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = subtitle.add_run(
        "BLIP Captions, Sentence-BERT Diversity, and Human vs AI Comparison"
    )
    run2.italic = True
    run2.font.size = Pt(14)
    run2.font.color.rgb = DARK

    document.add_paragraph()
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(
        "Project: Explainable Image Ambiguity Prediction Using Human and "
        "AI-Generated Caption Diversity with Computer Vision Features\n"
        "Modules: BlipCaptionGenerator, CaptionDiversityAnalyzer, "
        "MLDatasetBuilder\n"
        "Outputs: dataset/ai_dataset.csv, human vs AI diversity comparison"
    )
    meta_run.font.size = Pt(11)
    meta_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    document.add_page_break()


def add_heading(document: Document, text: str, level: int = 1) -> None:
    heading = document.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = ACCENT


def add_body(document: Document, text: str) -> None:
    p = document.add_paragraph(text)
    p.paragraph_format.space_after = Pt(8)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_numbered(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Number")


def add_code(document: Document, code: str) -> None:
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    for i, line in enumerate(code.strip("\n").split("\n")):
        run = p.add_run(("\n" if i > 0 else "") + line)
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x0A, 0x0A, 0x0A)
    p.paragraph_format.space_after = Pt(10)


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    document.add_paragraph()


def build() -> None:
    document = Document()
    set_base_font(document)
    add_title_page(document)

    add_heading(document, "1. Purpose")
    add_body(
        document,
        "Human COCO captions already support a caption-diversity feature. "
        "The AI Caption Diversity pipeline repeats the same measurement "
        "using BLIP-generated captions, so the project can compare how much "
        "AI descriptions disagree with each other versus how much human "
        "descriptions disagree.",
    )
    add_body(
        document,
        "Higher caption diversity means lower average semantic agreement "
        "among captions for the same image, which is treated as a signal "
        "of visual ambiguity.",
    )

    add_heading(document, "2. Pipeline Overview")
    add_numbered(
        document,
        [
            "Select the same image ids used in dataset/human_dataset.csv "
            "(typically 300 images, seed 42).",
            "Generate multiple AI captions per image with BLIP using "
            "beam search, top-k sampling, and nucleus sampling.",
            "Embed those captions with Sentence-BERT "
            "(all-MiniLM-L6-v2).",
            "Compute pairwise cosine similarity and "
            "Caption Diversity = 1 - Average Similarity.",
            "Merge with OpenCV image features into dataset/ai_dataset.csv.",
            "Compare mean caption_diversity against human_dataset.csv.",
        ],
    )

    add_heading(document, "3. Tools Used")
    add_table(
        document,
        ["Tool", "Role"],
        [
            ["BLIP (Salesforce/blip-image-captioning-base)", "AI caption generation"],
            ["Hugging Face Transformers + PyTorch", "Load and run BLIP"],
            ["Sentence-BERT (all-MiniLM-L6-v2)", "Caption embeddings"],
            ["CaptionDiversityAnalyzer", "Pairwise similarity + diversity score"],
            ["OpenCVFeatureExtractor", "Edge/entropy/brightness/contrast/texture"],
            ["MLDatasetBuilder", "Assemble tabular CSV rows"],
            ["Pillow", "Load RGB images from disk"],
            ["pandas", "CSV dataset I/O and mean comparison"],
        ],
    )

    add_heading(document, "4. BLIP Decoding Strategies")
    add_body(
        document,
        "Multiple captions are generated intentionally so diversity is "
        "meaningful. Strategy outputs are flattened and de-duplicated.",
    )
    add_table(
        document,
        ["Strategy", "Method", "Typical settings"],
        [
            ["beam_search", "Deterministic beam search", "num_beams >= return count"],
            ["top_k", "Top-k sampling", "do_sample=True, top_k=50"],
            ["nucleus", "Top-p (nucleus) sampling", "do_sample=True, top_p=0.9"],
        ],
    )

    add_heading(document, "5. Diversity Formula")
    add_code(
        document,
        """
Average Similarity = mean(pairwise cosine similarities)
Caption Diversity  = 1 - Average Similarity
""",
    )
    add_body(
        document,
        "The same formula is used for human COCO captions and BLIP "
        "captions, which makes the comparison fair.",
    )

    add_heading(document, "6. Dataset Columns")
    add_body(
        document,
        "ai_dataset.csv matches the human dataset feature schema "
        "(without ambiguity labels unless added later):",
    )
    add_bullets(
        document,
        [
            "image_id",
            "average_similarity, minimum_similarity, maximum_similarity, "
            "std_similarity, caption_diversity",
            "edge_density, entropy, brightness, contrast, color_variance, texture",
        ],
    )

    add_heading(document, "7. How It Is Used in This Project")
    add_bullets(
        document,
        [
            "Provides AI-side diversity features for ambiguity research.",
            "Supports human-vs-AI disagreement analysis.",
            "Feeds the FastAPI inference path when captions are generated "
            "by BLIP for an uploaded image.",
            "Complements OpenCV features in the supervised ambiguity model.",
        ],
    )

    add_heading(document, "8. Usage")
    add_body(document, "Generate / refresh the AI dataset (when the builder CLI is present):")
    add_code(
        document,
        """
python src/create_ai_dataset.py
# optional:
python src/create_ai_dataset.py --device cpu --num-return-sequences 2
""",
    )
    add_body(document, "Compare mean diversity against the human dataset:")
    add_code(
        document,
        """
python src/compare.py
""",
    )
    add_body(document, "Example comparison output format:")
    add_code(
        document,
        """
Human Diversity

0.38

AI Diversity

0.48
""",
    )
    add_body(
        document,
        "Exact numbers depend on the sampled images and BLIP captions. "
        "In this project’s 300-image run, AI captions were more diverse "
        "than human captions on average.",
    )

    add_heading(document, "9. Single-Image BLIP Demo")
    add_code(
        document,
        """
python src/blip_caption.py
python src/blip_caption.py --image-id 397133
""",
    )
    add_body(
        document,
        "JSON results are saved under results/captions/ "
        "(for example blip_397133.json), including generated captions "
        "and comparison against COCO references when available.",
    )

    add_heading(document, "10. Key Source Files")
    add_table(
        document,
        ["Path", "Responsibility"],
        [
            [
                "src/image_ambiguity/features/blip_captions.py",
                "BLIP generation + COCO comparison helpers",
            ],
            [
                "src/image_ambiguity/features/caption_diversity.py",
                "Diversity metrics from embeddings",
            ],
            [
                "src/image_ambiguity/features/sentence_embeddings.py",
                "Sentence-BERT encoding",
            ],
            [
                "src/image_ambiguity/pipeline/dataset_builder.py",
                "Build CSV rows (supports injected AI captions)",
            ],
            [
                "dataset/ai_dataset.csv",
                "AI caption-diversity + CV feature table",
            ],
            [
                "dataset/human_dataset.csv",
                "Human caption-diversity + CV feature table",
            ],
        ],
    )

    add_heading(document, "11. References")
    add_bullets(
        document,
        [
            "Li, J. et al. (2022). BLIP: Bootstrapping Language-Image "
            "Pre-training. ICML 2022.",
            "Reimers, N., & Gurevych, I. (2019). Sentence-BERT. EMNLP-IJCNLP.",
            "Lin, T.-Y. et al. (2014). Microsoft COCO. ECCV 2014.",
        ],
    )

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")


if __name__ == "__main__":
    build()
