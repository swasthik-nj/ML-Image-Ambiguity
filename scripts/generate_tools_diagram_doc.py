"""Generate a conference-ready doc explaining every tool/module used.

Draws the tool-stack diagram (via generate_tools_diagram.build_diagram) and
embeds it in a Word document alongside a tool-by-tool explanation, grouped by
pipeline stage, plus the rationale for each technology choice.

Run:

    python scripts/generate_tools_diagram_doc.py

Produces:
    results/figures/tools_architecture_diagram.png
    myDocs/Tools_and_Architecture_Diagram.docx
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).resolve().parent))

from generate_tools_diagram import build_diagram  # noqa: E402

OUTPUT_PATH = PROJECT_ROOT / "myDocs" / "Tools_and_Architecture_Diagram.docx"

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
DARK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x55, 0x55, 0x55)


def set_base_font(document: Document, name: str = "Calibri") -> None:
    style = document.styles["Normal"]
    style.font.name = name
    style.font.size = Pt(11)
    rpr = style.element.get_or_add_rPr()
    rFonts = rpr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)


def add_title_page(document: Document) -> None:
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Tools & Architecture Diagram")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = ACCENT

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = subtitle.add_run(
        "Every Module, Library, and Framework Used \u2014 End-to-End Map"
    )
    run2.italic = True
    run2.font.size = Pt(14)
    run2.font.color.rgb = DARK

    document.add_paragraph()
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(
        "Project: Explainable Image Ambiguity Prediction Using Human and "
        "AI-Generated Caption Diversity with Computer Vision Features\n"
        "Document type: Technology stack & architecture reference"
    )
    meta_run.font.size = Pt(11)
    meta_run.font.color.rgb = MUTED
    document.add_page_break()


def add_heading(document: Document, text: str, level: int = 1) -> None:
    heading = document.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = ACCENT


def add_body(document: Document, text: str) -> None:
    p = document.add_paragraph(text)
    p.paragraph_format.space_after = Pt(8)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        p = document.add_paragraph(item, style="List Bullet")
        p.paragraph_format.space_after = Pt(3)


def add_numbered(document: Document, items: list[str]) -> None:
    for item in items:
        p = document.add_paragraph(item, style="List Number")
        p.paragraph_format.space_after = Pt(3)


def add_code(document: Document, code: str) -> None:
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    for i, line in enumerate(code.strip("\n").split("\n")):
        run = p.add_run(("\n" if i > 0 else "") + line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x0A, 0x0A, 0x0A)
    p.paragraph_format.space_after = Pt(10)


def add_image(document: Document, path: Path, width_in: float = 6.3) -> None:
    document.add_picture(str(path), width=Inches(width_in))
    last_paragraph = document.paragraphs[-1]
    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_tool_table(document: Document, rows: list[tuple[str, str, str]]) -> None:
    """rows: (module, library/tool, role)."""
    table = document.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    header = table.rows[0].cells
    header[0].text = "Module"
    header[1].text = "Library / Tool"
    header[2].text = "Role in the Pipeline"
    for module, library, role in rows:
        cells = table.add_row().cells
        cells[0].text = module
        cells[1].text = library
        cells[2].text = role
    document.add_paragraph().paragraph_format.space_after = Pt(6)


def build() -> Path:
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    diagram_path = build_diagram()

    document = Document()
    set_base_font(document)
    add_title_page(document)

    # ------------------------------------------------------------------
    add_heading(document, "1. Purpose of This Document")
    add_body(
        document,
        "This document gives a single visual map of every tool, library, and "
        "custom module used across the project \u2014 from raw COCO data to the "
        "deployed FastAPI + React application \u2014 followed by a tool-by-tool "
        "explanation of why each one was chosen and what it contributes to "
        "the final ambiguity prediction.",
    )

    # ------------------------------------------------------------------
    add_heading(document, "2. End-to-End Tool & Module Diagram")
    add_body(
        document,
        "Boxes are colour-coded by pipeline stage. Arrows show data/feature "
        "flow from the raw dataset through to the interactive UI.",
    )
    add_image(document, diagram_path)
    document.add_paragraph()
    add_bullets(
        document,
        [
            "Blue \u2014 data source and loading (COCO, pycocotools).",
            "Purple \u2014 NLP / embedding tools (Sentence-BERT, BLIP, cosine "
            "similarity).",
            "Orange \u2014 classical computer vision (OpenCV).",
            "Green \u2014 dataset assembly and labeling pipeline (pandas).",
            "Red \u2014 supervised modeling (Random Forest, XGBoost, joblib).",
            "Yellow \u2014 explainability (SHAP).",
            "Teal / Lavender \u2014 serving layer (FastAPI) and UI (React).",
        ],
    )

    # ------------------------------------------------------------------
    add_heading(document, "3. Data & Captioning Tools")
    add_tool_table(
        document,
        [
            (
                "CocoDatasetLoader\nsrc/image_ambiguity/data/coco_loader.py",
                "pycocotools, Pillow",
                "Indexes captions_val2017.json and dataset/val2017/ images; "
                "exposes get_image(), get_captions(), and random sampling.",
            ),
            (
                "SentenceEmbeddingGenerator\nfeatures/sentence_embeddings.py",
                "sentence-transformers (all-MiniLM-L6-v2), torch",
                "Converts each caption into a 384-dim dense vector; auto "
                "detects GPU/CPU; batches inference for speed.",
            ),
            (
                "BlipCaptionGenerator\nfeatures/blip_captions.py",
                "transformers (Salesforce/blip-image-captioning-base), torch",
                "Generates AI captions (beam / top-k / nucleus / prompted) "
                "used for ai_dataset.csv and BLIP-forced predictions.",
            ),
        ],
    )

    # ------------------------------------------------------------------
    add_heading(document, "4. Feature Extraction Tools")
    add_tool_table(
        document,
        [
            (
                "CaptionDiversityAnalyzer\nfeatures/caption_diversity.py",
                "scikit-learn (cosine_similarity), numpy",
                "Pairwise cosine similarity across caption embeddings \u2192 "
                "average / min / max / std similarity and diversity score.",
            ),
            (
                "OpenCVFeatureExtractor\nfeatures/cv_features.py",
                "opencv-python, numpy",
                "Extracts edge density, entropy, brightness, contrast, "
                "color variance, and texture from the raw image pixels.",
            ),
        ],
    )

    # ------------------------------------------------------------------
    add_heading(document, "5. Dataset Assembly & Labeling Tools")
    add_tool_table(
        document,
        [
            (
                "MLDatasetBuilder\npipeline/dataset_builder.py",
                "pandas",
                "Joins diversity metrics + OpenCV features into one row per "
                "image \u2192 human_dataset.csv / ai_dataset.csv.",
            ),
            (
                "AmbiguityLabelGenerator\npipeline/label_generator.py",
                "pandas, matplotlib",
                "Applies Low (<0.35) / Medium (0.35\u20130.65) / High (\u22650.65) "
                "rules to caption_diversity; produces stats and plots.",
            ),
            (
                "enrich_high_diversity.py\nsrc/enrich_high_diversity.py",
                "pandas, pycocotools",
                "Mines additional High-diversity COCO images to reduce class "
                "imbalance before training.",
            ),
        ],
    )

    # ------------------------------------------------------------------
    add_heading(document, "6. Modeling & Explainability Tools")
    add_tool_table(
        document,
        [
            (
                "ModelTrainer\nmodels/trainer.py",
                "scikit-learn (RandomForestClassifier), xgboost, joblib",
                "80/20 stratified split, cross-validation, GridSearchCV "
                "tuning, class balancing, and evaluation "
                "(Accuracy/Precision/Recall/F1/ROC-AUC).",
            ),
            (
                "SHAPExplainer\nexplainability/shap_explainer.py",
                "shap (TreeExplainer)",
                "Summary, waterfall, and force plots plus top feature "
                "importance to explain individual predictions.",
            ),
            (
                "numba_stub\nexplainability/numba_stub.py",
                "pure-Python compatibility shim",
                "Lets shap import on Windows machines where an Application "
                "Control policy blocks numba's native DLLs.",
            ),
        ],
    )

    # ------------------------------------------------------------------
    add_heading(document, "7. Serving & Application Tools")
    add_tool_table(
        document,
        [
            (
                "FastAPI Backend\nbackend/app/",
                "fastapi, uvicorn, pydantic / pydantic-settings",
                "REST API exposing /upload, /features, /predict, /explain, "
                "/compare, /health for the trained pipeline.",
            ),
            (
                "AmbiguityInferenceService\nbackend/app/services/inference.py",
                "pandas, PIL",
                "Orchestrates caption resolution (user / COCO human / BLIP), "
                "feature extraction, prediction, and SHAP explanation.",
            ),
            (
                "React Frontend \u2013 Ambiguity Lens\nfrontend/",
                "Vite, React, TypeScript, Tailwind CSS, Recharts, React "
                "Router",
                "Upload UI, prediction dashboard with charts, and a "
                "human-vs-AI caption diversity comparison page.",
            ),
        ],
    )

    # ------------------------------------------------------------------
    add_heading(document, "8. Configuration, Logging & Dev Tools")
    add_tool_table(
        document,
        [
            (
                "Settings\nimage_ambiguity/config.py",
                "pydantic-settings, PyYAML, python-dotenv",
                "Centralized, typed configuration from configs/default.yaml "
                "and .env (paths, model names, device, batch sizes).",
            ),
            (
                "Logging setup\nimage_ambiguity/logging_config.py",
                "logging (stdlib), RotatingFileHandler",
                "Structured console + rotating file logs across every "
                "module and script.",
            ),
            (
                "Test suite\ntests/unit/, tests/",
                "pytest, pytest-cov",
                "Unit tests for every module (loader, SBERT, diversity, "
                "OpenCV, labeling, trainer, SHAP, numba stub).",
            ),
            (
                "Lint\n(project-wide)",
                "ruff",
                "Static analysis / style enforcement across the src/ tree.",
            ),
        ],
    )

    # ------------------------------------------------------------------
    add_heading(document, "9. Why These Tools Were Chosen")
    add_bullets(
        document,
        [
            "pycocotools is the reference implementation for COCO indexing "
            "\u2014 avoids re-implementing annotation parsing.",
            "Sentence-BERT (all-MiniLM-L6-v2) gives strong sentence "
            "similarity quality at a small, fast, CPU-friendly model size.",
            "OpenCV is the de-facto standard for classical, interpretable "
            "image descriptors (edges, entropy, texture) \u2014 cheap to "
            "compute and easy to explain to a non-ML audience.",
            "Random Forest + XGBoost were chosen together: Random Forest is "
            "robust and interpretable via SHAP; XGBoost typically pushes "
            "accuracy higher on tabular data \u2014 the trainer keeps the best "
            "of the two automatically.",
            "SHAP was chosen over simpler feature-importance methods "
            "because it gives per-prediction, signed contributions that map "
            "directly to a human explanation (\"this feature pushed the "
            "prediction toward High\").",
            "FastAPI was chosen for automatic OpenAPI docs (/docs), async "
            "support, and pydantic-based request/response validation.",
            "React + Vite + Tailwind + Recharts gives a fast dev loop and "
            "ready-made charts for the prediction and comparison pages.",
        ],
    )

    # ------------------------------------------------------------------
    add_heading(document, "10. Reproduce the Full Pipeline")
    add_code(
        document,
        """
# from project root, with the virtual environment active
$env:PYTHONPATH="src;."

python src/create_dataset.py          # human_dataset.csv (SBERT + OpenCV)
python src/generate_labels.py         # + ambiguity_label column
python src/enrich_high_diversity.py --max-high 40 --high-min 0.65
python src/train.py --balance-classes # Random Forest + XGBoost -> joblib
python src/shap_analysis.py           # SHAP plots + explanation

python scripts/generate_tools_diagram.py   # regenerate this diagram
""",
    )

    # ------------------------------------------------------------------
    add_heading(document, "11. Talking Points (Conference / Viva)")
    add_numbered(
        document,
        [
            "\"We reuse pycocotools for correctness on COCO parsing, and "
            "build everything else \u2014 diversity scoring, CV features, "
            "labeling, training, explainability \u2014 as small, testable, "
            "single-responsibility classes.\"",
            "\"Ambiguity is operationalized as measurable caption "
            "disagreement (Sentence-BERT diversity), not visual clutter.\"",
            "\"Random Forest and XGBoost are trained on the same 11-feature "
            "vector; we keep whichever scores best on the held-out set.\"",
            "\"SHAP closes the loop: every prediction served by the API can "
            "be explained feature-by-feature, not just classified.\"",
            "\"The same feature pipeline powers offline training (CSV) and "
            "the live FastAPI /predict and /explain endpoints \u2014 no "
            "train/serve skew.\"",
        ],
    )

    # ------------------------------------------------------------------
    add_heading(document, "References")
    add_bullets(
        document,
        [
            "Lin, T.-Y. et al. (2014). Microsoft COCO: Common Objects in "
            "Context.",
            "Reimers, N., & Gurevych, I. (2019). Sentence-BERT: Sentence "
            "Embeddings using Siamese BERT-Networks.",
            "Li, J. et al. (2022). BLIP: Bootstrapping Language-Image "
            "Pre-training.",
            "Bradski, G. (2000). The OpenCV Library.",
            "Breiman, L. (2001). Random Forests.",
            "Chen, T., & Guestrin, C. (2016). XGBoost: A Scalable Tree "
            "Boosting System.",
            "Lundberg, S. M., & Lee, S.-I. (2017). A Unified Approach to "
            "Interpreting Model Predictions (SHAP).",
        ],
    )

    document.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
