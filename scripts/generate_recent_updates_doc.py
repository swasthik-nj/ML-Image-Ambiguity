"""Generate documentation for recent project updates.

Run:

    python scripts/generate_recent_updates_doc.py

Produces: myDocs/Recent_Updates_High_Ambiguity_and_COCO_Captions.docx
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = (
    PROJECT_ROOT
    / "myDocs"
    / "Recent_Updates_High_Ambiguity_and_COCO_Captions.docx"
)

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
DARK = RGBColor(0x22, 0x22, 0x22)


def set_base_font(document: Document, name: str = "Calibri") -> None:
    style = document.styles["Normal"]
    style.font.name = name
    style.font.size = Pt(11)
    rpr = style.element.get_or_add_rPr()
    rFonts = rpr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)


def add_title_page(document: Document) -> None:
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Recent Project Updates")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = ACCENT

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = subtitle.add_run(
        "High-Ambiguity Class Balancing, Retraining, and COCO Caption Fix"
    )
    run2.italic = True
    run2.font.size = Pt(14)
    run2.font.color.rgb = DARK

    document.add_paragraph()
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(
        "Project: Explainable Image Ambiguity Prediction\n"
        "Date: July 2026\n"
        "Focus: Keep High ≥ 0.65, fix training imbalance, "
        "and stop BLIP from mislabeling High COCO images as Medium"
    )
    meta_run.font.size = Pt(11)
    meta_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    document.add_page_break()


def add_heading(document: Document, text: str, level: int = 1) -> None:
    heading = document.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = ACCENT


def add_body(document: Document, text: str) -> None:
    p = document.add_paragraph(text)
    p.paragraph_format.space_after = Pt(8)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        p = document.add_paragraph(item, style="List Bullet")
        p.paragraph_format.space_after = Pt(4)


def add_code(document: Document, text: str) -> None:
    p = document.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    p.paragraph_format.space_after = Pt(10)


def build() -> Path:
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    set_base_font(document)
    add_title_page(document)

    add_heading(document, "1. Summary")
    add_body(
        document,
        "This update keeps the research-correct ambiguity thresholds "
        "(Low < 0.35, Medium 0.35–0.65, High ≥ 0.65) while fixing two "
        "practical problems: (1) almost no High samples in training, and "
        "(2) High COCO images scoring Medium in the UI when BLIP captions "
        "were used instead of human captions.",
    )
    add_bullets(
        document,
        [
            "Enriched human_dataset.csv with more High-diversity COCO images "
            "(High: 4 → 34).",
            "Retrained Random Forest / XGBoost with train-split class balancing.",
            "Prediction now auto-loads COCO human captions for val2017 filenames.",
            "BLIP sampling was made more diverse for true AI-caption runs.",
            "UI still shows the previous Prediction page design.",
        ],
    )

    add_heading(document, "2. Problem: Everything Looked Medium")
    add_heading(document, "2.1 Training imbalance", level=2)
    add_body(
        document,
        "Before enrichment, the labeled human set was heavily skewed toward "
        "Medium (about 187 Medium, 109 Low, only 4 High). Models therefore "
        "tended to predict Medium with very high confidence.",
    )
    add_heading(document, "2.2 BLIP vs human captions", level=2)
    add_body(
        document,
        "Images labeled High in the dataset are High because human COCO "
        "annotators disagree (caption diversity ≥ 0.65). BLIP often produces "
        "near-duplicate captions for the same image (for example many variants "
        "of “croissants”), so diversity falls around 0.40 and the label "
        "correctly becomes Medium for those AI captions. This is not a "
        "classifier crash; it is a caption-source mismatch.",
    )
    add_body(
        document,
        "Example: dataset/val2017/000000538236.jpg has human diversity ≈ 0.80 "
        "(High). With BLIP-only captions it scored diversity ≈ 0.40 (Medium).",
    )

    add_heading(document, "3. Fix A — Keep Thresholds, Balance Training")
    add_heading(document, "3.1 Enrich High images", level=2)
    add_body(
        document,
        "Script src/enrich_high_diversity.py scans unused COCO val2017 images, "
        "keeps those with caption_diversity ≥ 0.65, extracts OpenCV features, "
        "and appends them to dataset/human_dataset.csv. Thresholds stay "
        "unchanged.",
    )
    add_code(
        document,
        "python src/enrich_high_diversity.py --max-high 40 --high-min 0.65",
    )
    add_body(
        document,
        "Result after enrichment: 330 rows — Low 109, Medium 187, High 34.",
    )

    add_heading(document, "3.2 Train-time class balancing", level=2)
    add_body(
        document,
        "ModelTrainer (src/image_ambiguity/models/trainer.py) now oversamples "
        "minority classes in the training split only (not the test split). "
        "Random Forest also uses class_weight='balanced'. Enabled by default "
        "in src/train.py via --balance-classes / --no-balance-classes.",
    )
    add_code(
        document,
        "python src/train.py --balance-classes",
    )
    add_body(
        document,
        "Artifacts updated: models/best_model.joblib, models/random_forest.joblib, "
        "models/xgboost.joblib, results/metrics/training_metrics.json.",
    )

    add_heading(document, "4. Fix B — Use COCO Human Captions for Dataset Images")
    add_body(
        document,
        "When the uploaded filename looks like a COCO id (for example "
        "000000538236.jpg), the backend looks up the official human captions "
        "from captions_val2017.json and prefers them unless force_blip=true.",
    )
    add_bullets(
        document,
        [
            "backend/app/services/inference.py — parse_coco_image_id, "
            "lookup_coco_captions, resolve_captions, caption_source.",
            "Upload metadata now stores coco_image_id and coco_captions.",
            "API forms accept force_blip on /features, /predict, /explain.",
            "Responses include caption_source: user | coco_human | blip.",
            "Frontend auto-turns BLIP off for COCO filenames and fills human "
            "captions after upload.",
        ],
    )
    add_body(
        document,
        "Verified: 000000538236.jpg with COCO captions → diversity 0.798 → High.",
    )

    add_heading(document, "5. Fix C — More Diverse BLIP (When Forced)")
    add_body(
        document,
        "If the user explicitly enables BLIP (force_blip), captions are still "
        "AI-generated. To reduce near-duplicates, BLIP now uses higher "
        "temperature / more return sequences and additional prompt-conditioned "
        "generations (src/image_ambiguity/features/blip_captions.py).",
    )
    add_body(
        document,
        "Note: even with richer BLIP sampling, AI captions may still agree more "
        "than humans. That human–AI gap is a valid research comparison finding, "
        "not something that should be hidden by lowering the High threshold.",
    )

    add_heading(document, "6. How to Demo High Ambiguity Correctly")
    add_bullets(
        document,
        [
            "Restart API: $env:PYTHONPATH=\"src;.\"; uvicorn app:app --reload",
            "Start UI: cd frontend; npm run dev",
            "Upload dataset/val2017/000000538236.jpg (or another High id).",
            "Leave BLIP unchecked so COCO human captions load.",
            "Expect Predicted ambiguity = High and caption diversity ≈ 0.65–0.80.",
            "Optional: check BLIP to compare AI caption diversity for the same image.",
        ],
    )
    add_heading(document, "6.1 Suggested High images", level=2)
    add_bullets(
        document,
        [
            "000000538236.jpg — diversity ≈ 0.80",
            "000000546823.jpg — diversity ≈ 0.79",
            "000000323263.jpg — diversity ≈ 0.79",
            "000000263425.jpg — diversity ≈ 0.79",
            "000000060932.jpg — diversity ≈ 0.76",
        ],
    )

    add_heading(document, "7. Label Rules (Unchanged)")
    add_bullets(
        document,
        [
            "Low: caption_diversity < 0.35",
            "Medium: 0.35 ≤ caption_diversity < 0.65",
            "High: caption_diversity ≥ 0.65",
        ],
    )
    add_body(
        document,
        "Ambiguity means caption disagreement, not visual clutter. A busy "
        "collage can still be Medium if captions agree on the content.",
    )

    add_heading(document, "8. Key Files Touched")
    add_bullets(
        document,
        [
            "src/enrich_high_diversity.py",
            "src/train.py",
            "src/image_ambiguity/models/trainer.py",
            "src/image_ambiguity/features/blip_captions.py",
            "backend/app/services/inference.py",
            "backend/app/api/routes.py",
            "backend/app/schemas.py",
            "frontend/src/pages/PredictionPage.tsx",
            "frontend/src/api/client.ts",
            "frontend/src/types.ts",
            "dataset/human_dataset.csv",
            "models/best_model.joblib",
        ],
    )

    add_heading(document, "9. Research Takeaway")
    add_body(
        document,
        "Keeping High ≥ 0.65 preserves the thesis definition: High means "
        "captions strongly disagree. Balancing and enriching the High class "
        "makes the classifier fairer. Automatically using COCO human captions "
        "for known val images makes demos match the labeled dataset. Forcing "
        "BLIP intentionally measures AI caption agreement and may yield Medium "
        "even on human-High images — a useful human vs AI diversity result.",
    )

    document.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
