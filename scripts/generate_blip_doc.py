"""One-off script to generate the BLIP caption generation explanation document.

Run once with:

    python scripts/generate_blip_doc.py

Produces: myDocs/BLIP_Caption_Generation_Explanation.docx
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = PROJECT_ROOT / "myDocs" / "BLIP_Caption_Generation_Explanation.docx"

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
DARK = RGBColor(0x22, 0x22, 0x22)


def set_base_font(document: Document, name: str = "Calibri") -> None:
    style = document.styles["Normal"]
    style.font.name = name
    style.font.size = Pt(11)
    rpr = style.element.get_or_add_rPr()
    rFonts = rpr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)


def add_title_page(document: Document) -> None:
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("BLIP AI Caption Generation")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = ACCENT

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = subtitle.add_run(
        "Implementation, Tools, Project Role, and Usage Guide"
    )
    run2.italic = True
    run2.font.size = Pt(15)
    run2.font.color.rgb = DARK

    document.add_paragraph()
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(
        "Project: Explainable Image Ambiguity Prediction Using Human and "
        "AI-Generated Caption Diversity with Computer Vision Features\n"
        "Module: image_ambiguity.features.blip_captions.BlipCaptionGenerator\n"
        "CLI: src/blip_caption.py\n"
        "Model: Salesforce/blip-image-captioning-base"
    )
    meta_run.font.size = Pt(11)
    meta_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    document.add_page_break()


def add_heading(document: Document, text: str, level: int = 1) -> None:
    heading = document.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = ACCENT


def add_body(document: Document, text: str) -> None:
    p = document.add_paragraph(text)
    p.paragraph_format.space_after = Pt(8)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_numbered(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Number")


def add_code(document: Document, code: str) -> None:
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    for i, line in enumerate(code.strip("\n").split("\n")):
        run = p.add_run(("\n" if i > 0 else "") + line)
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x0A, 0x0A, 0x0A)
    p.paragraph_format.space_after = Pt(10)


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    document.add_paragraph()


def build() -> None:
    document = Document()
    set_base_font(document)
    add_title_page(document)

    # 1. Purpose
    add_heading(document, "1. Purpose of the BLIP Module")
    add_body(
        document,
        "This project studies image ambiguity by measuring how much "
        "captions disagree. Human captions already come from COCO. The "
        "BLIP module adds AI-generated captions for the same image, so "
        "human-vs-AI and AI-vs-AI disagreement can later be used as "
        "ambiguity features.",
    )
    add_body(
        document,
        "Given one image, the module generates multiple captions with "
        "three decoding strategies (beam search, top-k sampling, and "
        "nucleus sampling), compares them with COCO human captions, "
        "saves a JSON record, and prints the result.",
    )

    # 2. What is BLIP
    add_heading(document, "2. What is BLIP?")
    add_body(
        document,
        "BLIP (Bootstrapping Language-Image Pre-training) is a "
        "vision-language model from Salesforce Research. The variant used "
        "here is an image-captioning model: it takes an image and "
        "produces a natural-language description.",
    )
    add_bullets(
        document,
        [
            "Model id: Salesforce/blip-image-captioning-base",
            "Source: Hugging Face Hub (downloaded on first run)",
            "Task: conditional text generation from image features",
            "Library API: transformers.BlipProcessor + "
            "transformers.BlipForConditionalGeneration",
        ],
    )

    # 3. Tools
    add_heading(document, "3. Tools Used")
    add_body(
        document,
        "The following libraries and project components are used by the "
        "BLIP implementation:",
    )
    add_table(
        document,
        ["Tool / Component", "Role in this module"],
        [
            [
                "PyTorch (torch)",
                "Runs the BLIP neural network on CPU or CUDA",
            ],
            [
                "Hugging Face Transformers",
                "Loads BlipProcessor and BlipForConditionalGeneration",
            ],
            [
                "Pillow (PIL)",
                "Opens and converts input images to RGB",
            ],
            [
                "pycocotools + CocoDatasetLoader",
                "Loads COCO images and human reference captions",
            ],
            [
                "image_ambiguity.config",
                "Resolves dataset paths, device, and results directories",
            ],
            [
                "image_ambiguity.utils.common",
                "save_json / ensure_dir / timed helpers",
            ],
            [
                "argparse CLI (src/blip_caption.py)",
                "Runnable entrypoint for demos and experiments",
            ],
        ],
    )
    add_body(
        document,
        "These dependencies are already listed in requirements.txt "
        "(transformers, torch, Pillow, pycocotools). No extra BLIP-only "
        "package is required beyond Hugging Face Transformers.",
    )

    # 4. Project role
    add_heading(document, "4. How BLIP Fits in This Project")
    add_body(
        document,
        "The research pipeline separates human captions, AI captions, "
        "diversity scoring, CV features, labeling, and supervised "
        "training. BLIP sits in the AI-caption generation stage:",
    )
    add_numbered(
        document,
        [
            "Load a COCO image (dataset/val2017) and its human captions "
            "(dataset/annotations/captions_val2017.json).",
            "Run BLIP with multiple decoding strategies to produce "
            "AI captions for that image.",
            "Compare AI captions with COCO captions (exact match + "
            "word-overlap / Jaccard).",
            "Save results under results/captions/ for later diversity "
            "analysis (Sentence-BERT) and ambiguity modeling.",
        ],
    )
    add_body(
        document,
        "In short: COCO provides human captions; BLIP provides AI "
        "captions; caption diversity modules measure disagreement; the "
        "ML trainer predicts Low / Medium / High ambiguity.",
    )

    # 5. Implementation
    add_heading(document, "5. Implementation in This Repository")
    add_heading(document, "5.1 Source files", level=2)
    add_table(
        document,
        ["Path", "Responsibility"],
        [
            [
                "src/image_ambiguity/features/blip_captions.py",
                "Core library: BlipCaptionGenerator, comparison helpers, "
                "JSON serialization",
            ],
            [
                "src/blip_caption.py",
                "Thin CLI that loads COCO, runs BLIP, prints and saves JSON",
            ],
            [
                "tests/unit/test_blip_captions.py",
                "Unit tests for comparison helpers and generator API",
            ],
            [
                "results/captions/blip_<image_id>.json",
                "Default output location for saved captions",
            ],
        ],
    )

    add_heading(document, "5.2 Main class API", level=2)
    add_body(
        document,
        "BlipCaptionGenerator wraps model loading and the three decoding "
        "strategies. Typical library usage:",
    )
    add_code(
        document,
        """
from PIL import Image
from image_ambiguity.features.blip_captions import BlipCaptionGenerator

generator = BlipCaptionGenerator(device="auto")
generator.load_model()

image = Image.open("dataset/val2017/000000397133.jpg").convert("RGB")
result = generator.caption_image(
    image,
    image_id=397133,
    file_name="000000397133.jpg",
    coco_captions=["A man is in a kitchen making pizzas.", ...],
)
generator.save_result(result, "results/captions/blip_397133.json")
payload = result.to_dict()  # JSON-serializable dict
""",
    )

    add_heading(document, "5.3 Decoding strategies", level=2)
    add_body(
        document,
        "Multiple captions are generated intentionally. Different "
        "decoding strategies explore different parts of the model's "
        "output distribution, which is useful when studying caption "
        "disagreement as an ambiguity signal.",
    )
    add_table(
        document,
        ["Strategy key", "Method", "Default settings"],
        [
            [
                "beam_search",
                "Deterministic beam search",
                "num_beams=5, early_stopping=True",
            ],
            [
                "top_k",
                "Top-k sampling (stochastic)",
                "do_sample=True, top_k=50, temperature=1.0",
            ],
            [
                "nucleus",
                "Nucleus / top-p sampling",
                "do_sample=True, top_p=0.9, temperature=1.0",
            ],
        ],
    )
    add_body(
        document,
        "By default each strategy returns up to 3 unique captions "
        "(num_return_sequences=3). Duplicate decoded strings are removed "
        "while preserving order.",
    )

    add_heading(document, "5.4 Comparison with COCO captions", level=2)
    add_body(
        document,
        "After generation, compare_with_coco() builds a structured "
        "comparison block:",
    )
    add_bullets(
        document,
        [
            "exact_matches — generated captions that match a COCO "
            "caption after normalization (lowercase, punctuation stripped)",
            "best_overlaps — for each generated caption, the COCO caption "
            "with the highest word-set Jaccard similarity",
            "n_coco / n_generated — counts for quick inspection",
        ],
    )

    # 6. Usage
    add_heading(document, "6. Usage")
    add_heading(document, "6.1 Prerequisites", level=2)
    add_numbered(
        document,
        [
            "Install dependencies: pip install -r requirements.txt",
            "Install the package editable (optional but recommended): "
            "pip install -e .",
            "Ensure COCO files exist: "
            "dataset/annotations/captions_val2017.json and "
            "dataset/val2017/*.jpg",
        ],
    )

    add_heading(document, "6.2 CLI commands", level=2)
    add_body(document, "From the project root:")
    add_code(
        document,
        """
# Default demo (image_id=397133)
python src/blip_caption.py

# Choose another COCO image
python src/blip_caption.py --image-id 397133

# Caption a direct image path
python src/blip_caption.py --image dataset/val2017/000000397133.jpg

# Force CPU / CUDA and control how many captions per strategy
python src/blip_caption.py --device cpu --num-return-sequences 2
python src/blip_caption.py --device cuda --num-return-sequences 3

# Custom output path / skip COCO comparison
python src/blip_caption.py --output results/captions/demo.json
python src/blip_caption.py --no-compare
""",
    )

    add_heading(document, "6.3 Important CLI flags", level=2)
    add_table(
        document,
        ["Flag", "Meaning"],
        [
            ["--image-id", "COCO image id (default: 397133)"],
            ["--image", "Direct path to an image file"],
            ["--annotation-file", "Override captions_*.json path"],
            ["--image-dir", "Override image directory"],
            ["--model", "Hugging Face BLIP model id"],
            ["--device", "cpu | cuda | auto"],
            ["--num-return-sequences", "Captions per strategy (default: 3)"],
            ["--output", "JSON output path"],
            ["--no-compare", "Do not load/compare COCO captions"],
        ],
    )

    add_heading(document, "6.4 First-run note", level=2)
    add_body(
        document,
        "The first execution downloads Salesforce/blip-image-captioning-base "
        "from Hugging Face into the local cache. Later runs reuse the "
        "cached weights and are much faster. On machines without a GPU, "
        "use --device cpu.",
    )

    # 7. Output JSON
    add_heading(document, "7. Output JSON Format")
    add_body(
        document,
        "Default save path: results/captions/blip_<image_id>.json. "
        "Example structure:",
    )
    add_code(
        document,
        """
{
  "image_id": 397133,
  "file_name": "000000397133.jpg",
  "model": "Salesforce/blip-image-captioning-base",
  "device": "cpu",
  "generated_captions": {
    "beam_search": ["a woman standing in a kitchen", "..."],
    "top_k": ["pots are on walls", "..."],
    "nucleus": ["a kitchen with a wood counter", "..."]
  },
  "coco_captions": [
    "A man is in a kitchen making pizzas.",
    "..."
  ],
  "comparison": {
    "n_coco": 5,
    "n_generated": 6,
    "exact_matches": [],
    "best_overlaps": {
      "beam_search": [
        {
          "generated": "a woman standing in a kitchen",
          "best_coco": "A person standing by a stove in a kitchen.",
          "jaccard": 0.5
        }
      ]
    }
  }
}
""",
    )

    # 8. Example result
    add_heading(document, "8. Example Result (Image 397133)")
    add_body(
        document,
        "On the default validation image (a kitchen scene), BLIP "
        "produced short scene-level captions such as “a woman standing "
        "in a kitchen”, while COCO humans described baking / pizza "
        "preparation in more detail. Exact string matches were rare, "
        "but word-overlap scores showed partial agreement "
        "(for example Jaccard ≈ 0.5 against “A person standing by a "
        "stove in a kitchen.”). This human–AI divergence is exactly the "
        "kind of signal the larger ambiguity pipeline is designed to "
        "quantify.",
    )

    # 9. Testing
    add_heading(document, "9. Testing")
    add_body(document, "Unit tests (mocked model, no network download):")
    add_code(document, "python -m pytest tests/unit/test_blip_captions.py -q")
    add_body(document, "End-to-end smoke test (downloads model on first run):")
    add_code(document, "python src/blip_caption.py")

    # 10. Talking points
    add_heading(document, "10. Conference / Report Talking Points")
    add_bullets(
        document,
        [
            "We generate AI captions with BLIP "
            "(Salesforce/blip-image-captioning-base) via Hugging Face "
            "Transformers and PyTorch.",
            "Three decoding strategies — beam search, top-k, and nucleus "
            "sampling — produce multiple AI captions per image.",
            "Generated captions are compared with COCO human captions "
            "using exact match and word-level Jaccard overlap.",
            "Results are persisted as JSON under results/captions/ for "
            "reproducible downstream diversity analysis.",
            "In the full system, these AI captions complement human COCO "
            "captions so ambiguity can be studied from caption "
            "disagreement, not from a single description alone.",
        ],
    )

    # 11. References
    add_heading(document, "11. References")
    add_bullets(
        document,
        [
            "Li, J., Li, D., Xiong, C., & Hoi, S. (2022). BLIP: "
            "Bootstrapping Language-Image Pre-training for Unified "
            "Vision-Language Understanding and Generation. ICML 2022.",
            "Salesforce BLIP image captioning model on Hugging Face: "
            "Salesforce/blip-image-captioning-base",
            "Wolf, T. et al. (2020). Transformers: State-of-the-Art "
            "Natural Language Processing. EMNLP 2020 (Hugging Face).",
            "Lin, T.-Y. et al. (2014). Microsoft COCO: Common Objects in "
            "Context. ECCV 2014.",
        ],
    )

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")


if __name__ == "__main__":
    build()
