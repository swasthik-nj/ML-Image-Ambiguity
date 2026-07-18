"""One-off script to generate the Dataset Creation explanation document.

Run once with:

    python scripts/generate_dataset_creation_doc.py

Produces: myDocs/Dataset_Creation_Module_Explanation.docx
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = PROJECT_ROOT / "myDocs" / "Dataset_Creation_Module_Explanation.docx"

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
DARK = RGBColor(0x22, 0x22, 0x22)


def set_base_font(document: Document, name: str = "Calibri") -> None:
    style = document.styles["Normal"]
    style.font.name = name
    style.font.size = Pt(11)
    rpr = style.element.get_or_add_rPr()
    rFonts = rpr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)


def add_title_page(document: Document) -> None:
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Building the Merged Ambiguity Dataset")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = ACCENT

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = subtitle.add_run(
        "Fusing Caption Diversity and OpenCV Features into One ML-Ready Table"
    )
    run2.italic = True
    run2.font.size = Pt(15)
    run2.font.color.rgb = DARK

    document.add_paragraph()
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(
        "Project: Explainable Image Ambiguity Prediction Using Human and "
        "AI-Generated Caption Diversity with Computer Vision Features\n"
        "Module: image_ambiguity.pipeline.dataset_builder.MLDatasetBuilder\n"
        "Output: dataset/human_dataset.csv"
    )
    meta_run.font.size = Pt(11)
    meta_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    document.add_page_break()


def add_heading(document: Document, text: str, level: int = 1) -> None:
    heading = document.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = ACCENT


def add_body(document: Document, text: str) -> None:
    p = document.add_paragraph(text)
    p.paragraph_format.space_after = Pt(8)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_numbered(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Number")


def add_code(document: Document, code: str) -> None:
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    for i, line in enumerate(code.strip("\n").split("\n")):
        run = p.add_run(("\n" if i > 0 else "") + line)
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x0A, 0x0A, 0x0A)
    p.paragraph_format.space_after = Pt(10)


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    document.add_paragraph()


def build() -> None:
    document = Document()
    set_base_font(document)
    add_title_page(document)

    # 1. Motivation
    add_heading(document, "1. Motivation")
    add_body(
        document,
        "The two feature families built so far in this project -- "
        "caption diversity (language-based) and OpenCV features "
        "(vision-based) -- each capture a different, partial view of "
        "image ambiguity. Neither is trained on labels; both are purely "
        "descriptive statistics. To actually train and evaluate an "
        "ambiguity prediction model, these two feature families must be "
        "fused into a single, consistent, per-image feature table."
    )
    add_body(
        document,
        "The Dataset Creation module is the fusion step: it takes a "
        "sample of COCO images, computes both feature families for each "
        "one, merges them into a single row, and writes the result to a "
        "reusable CSV -- the human_dataset.csv referenced throughout the "
        "rest of the pipeline (so named because it is built from human "
        "COCO captions, as distinct from a future AI-caption variant)."
    )

    # 2. What it produces
    add_heading(document, "2. What the Module Produces")
    add_body(
        document,
        "MLDatasetBuilder.build() returns a pandas DataFrame with exactly "
        "one row per image and twelve columns spanning both feature "
        "families:"
    )
    add_table(
        document,
        ["Column", "Source", "Meaning"],
        [
            ["image_id", "COCO annotation", "Unique identifier for the image"],
            ["average_similarity", "Caption diversity", "Mean pairwise cosine similarity across captions"],
            ["minimum_similarity", "Caption diversity", "Most disagreeing caption pair"],
            ["maximum_similarity", "Caption diversity", "Most agreeing caption pair"],
            ["std_similarity", "Caption diversity", "Spread of agreement across caption pairs"],
            ["caption_diversity", "Caption diversity", "1 - average_similarity"],
            ["edge_density", "OpenCV", "Fraction of Canny edge pixels"],
            ["entropy", "OpenCV", "Shannon entropy of intensity histogram"],
            ["brightness", "OpenCV", "Mean grayscale intensity"],
            ["contrast", "OpenCV", "Standard deviation of grayscale intensity"],
            ["color_variance", "OpenCV", "Mean per-channel pixel variance"],
            ["texture", "OpenCV", "Laplacian variance (texture energy)"],
        ],
    )

    # 3. Pipeline
    add_heading(document, "3. How the Dataset Is Assembled")
    add_numbered(
        document,
        [
            "Sample image ids: 300 image ids are drawn from "
            "captions_val2017.json using a seeded random sample "
            "(reproducible across runs).",
            "Collect captions: for each sampled image, its human caption "
            "set is retrieved via CocoDatasetLoader.",
            "Batch-encode once: every caption from every sampled image is "
            "flattened into a single list and encoded in one "
            "SentenceEmbeddingGenerator call, rather than one model call "
            "per image -- turning ~300 small calls into a single large, "
            "efficient batch.",
            "Slice and score: each image's embedding slice is passed to "
            "CaptionDiversityAnalyzer.compute() to recover its five "
            "diversity statistics.",
            "Extract CV features: OpenCVFeatureExtractor.extract() runs "
            "independently on each image's JPEG file to obtain the six "
            "vision-based statistics.",
            "Merge and impute: all twelve values are combined into one "
            "row per image; missing values (e.g. an image with fewer "
            "than two captions) are imputed rather than dropped.",
            "Persist: the final DataFrame is written to "
            "dataset/human_dataset.csv via pandas.",
        ],
    )

    # 4. Why batch encoding matters
    add_heading(document, "4. Why Batch Encoding Matters Here")
    add_body(
        document,
        "A naive implementation would call generate_embeddings() once "
        "per image (300 separate model invocations for ~1,500 total "
        "captions). Instead, MLDatasetBuilder flattens all captions "
        "across all sampled images into one list and calls the encoder "
        "exactly once, letting Sentence-BERT batch efficiently across the "
        "full caption pool before the embeddings are sliced back apart "
        "per image."
    )
    add_code(
        document,
        """
# Conceptually:
flat_captions = [c for image in images for c in image.captions]
all_embeddings = embedding_generator.generate_embeddings(flat_captions)
# then slice all_embeddings[start:end] back out per image
""",
    )
    add_body(
        document,
        "In this codebase, building the full 300-image dataset -- after "
        "the model was already loaded -- completed in under 6 seconds, "
        "demonstrating the practical benefit of batching for dataset-scale "
        "feature generation."
    )

    # 5. Handling missing values
    add_heading(document, "5. Handling Missing Values")
    add_body(
        document,
        "Two failure modes are anticipated and handled explicitly rather "
        "than allowed to crash the pipeline or silently drop rows:"
    )
    add_bullets(
        document,
        [
            "Fewer than two captions for an image: pairwise cosine "
            "similarity is undefined with a single caption, so all five "
            "diversity columns are set to NaN for that row.",
            "An unreadable or missing image file: OpenCV feature "
            "extraction fails gracefully, and all six CV columns are set "
            "to NaN for that row, with the error logged.",
        ],
    )
    add_body(
        document,
        "After every row is built, handle_missing_values() imputes any "
        "remaining NaN cells with that column's median across the sampled "
        "dataset (falling back to 0.0 only if an entire column is "
        "missing). Critically, no row is ever dropped -- every sampled "
        "image_id is preserved in the final table, which keeps the "
        "dataset size predictable and avoids silently shrinking the "
        "sample."
    )

    # 6. Worked example
    add_heading(document, "6. Worked Example (from this codebase)")
    add_body(
        document,
        "Running python src/create_dataset.py against COCO val2017 "
        "produced the following result:"
    )
    add_table(
        document,
        ["Metric", "Value"],
        [
            ["Images sampled", "300"],
            ["Rows written", "300"],
            ["Output file", "dataset/human_dataset.csv"],
            ["Columns", "12 (image_id + 5 diversity + 6 CV features)"],
            ["Missing values after impute", "0"],
            ["Build time (model pre-loaded)", "~5.6 seconds for 300 images"],
        ],
    )
    add_body(
        document,
        "A sample row (image_id 301061) shows average_similarity 0.74, "
        "caption_diversity 0.26 (low ambiguity signal from captions), "
        "alongside edge_density 0.059, entropy 7.76, and brightness "
        "113.4 -- illustrating how the two feature families sit "
        "side by side in a single, model-ready row."
    )

    # 7. Implementation
    add_heading(document, "7. Implementation in This Codebase")
    add_table(
        document,
        ["Method", "Responsibility"],
        [
            ["select_image_ids(sample_size, seed)", "Reproducibly samples COCO image ids."],
            ["build(image_ids=None, sample_size=300, seed=42)", "Orchestrates caption collection, batch encoding, diversity + CV extraction, and imputation; returns a DataFrame."],
            ["handle_missing_values(df)", "Static method: median-imputes NaNs per column without dropping rows."],
            ["save_csv(df, path)", "Writes the final DataFrame to CSV, creating parent directories as needed."],
        ],
    )
    add_code(
        document,
        """
builder = MLDatasetBuilder(
    annotation_file="dataset/annotations/captions_val2017.json",
    image_dir="dataset/val2017",
)
df = builder.build(sample_size=300, seed=42)
builder.save_csv(df, "dataset/human_dataset.csv")

print(len(df))  # 300
""",
    )
    add_body(
        document,
        "MLDatasetBuilder accepts injected SentenceEmbeddingGenerator, "
        "CaptionDiversityAnalyzer, and OpenCVFeatureExtractor instances, "
        "so the same class can be reused with different models, devices, "
        "or thresholds -- and can be unit tested end-to-end with mocked "
        "components instead of loading real neural network weights."
    )

    # 8. Why explainable
    add_heading(document, "8. Why This Design Supports Explainability")
    add_body(
        document,
        "Because every column in human_dataset.csv traces back to a "
        "specific, documented computation -- cosine similarity for "
        "captions, or a concrete OpenCV operation for pixels -- any "
        "downstream model trained on this table inherits fully "
        "interpretable inputs. A SHAP explanation citing 'caption_"
        "diversity' or 'texture' as a top contributor can be immediately "
        "traced back through this module to the exact captions or image "
        "regions responsible, closing the loop between raw evidence and "
        "final prediction."
    )

    # 9. Talking points
    add_heading(document, "9. Conference Talking Points (Summary)")
    add_bullets(
        document,
        [
            "MLDatasetBuilder fuses two independently interpretable "
            "feature families -- Sentence-BERT caption diversity and "
            "classical OpenCV image statistics -- into one 12-column "
            "table, one row per image.",
            "Captions across all sampled images are batch-encoded in a "
            "single Sentence-BERT call, building a 300-image dataset in "
            "under 6 seconds once the model is loaded.",
            "Missing data (single-caption images, unreadable files) is "
            "handled explicitly: affected columns are set to NaN, then "
            "imputed with the column median -- no rows are ever dropped.",
            "The result, human_dataset.csv, is a reproducible (seeded), "
            "versionable artifact that decouples feature engineering from "
            "model training.",
            "Every feature's provenance is documented and traceable, "
            "keeping the eventual ambiguity model's explanations grounded "
            "in concrete, human-verifiable evidence.",
        ],
    )

    # 10. References
    add_heading(document, "10. References")
    add_bullets(
        document,
        [
            "Lin, T.-Y. et al. (2014). Microsoft COCO: Common Objects in "
            "Context. ECCV 2014.",
            "Reimers, N., & Gurevych, I. (2019). Sentence-BERT: Sentence "
            "Embeddings using Siamese BERT-Networks. EMNLP-IJCNLP 2019.",
            "Bradski, G. (2000). The OpenCV Library. Dr. Dobb's Journal "
            "of Software Tools.",
            "McKinney, W. (2010). Data Structures for Statistical "
            "Computing in Python (pandas). SciPy 2010.",
        ],
    )

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")


if __name__ == "__main__":
    build()
